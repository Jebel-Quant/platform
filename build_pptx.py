from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import copy

# Colour palette
NAVY       = RGBColor(0x0D, 0x1B, 0x2A)
BLUE       = RGBColor(0x1E, 0x40, 0xAF)
ACCENT     = RGBColor(0x60, 0xA5, 0xFA)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xE5, 0xE7, 0xEB)
MID_GRAY   = RGBColor(0x6B, 0x72, 0x80)

W = Inches(13.33)
H = Inches(7.5)


def new_prs():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H
    return prs


def blank_slide(prs):
    blank = prs.slide_layouts[6]
    return prs.slides.add_slide(blank)


def fill_solid(shape, colour):
    shape.fill.solid()
    shape.fill.fore_color.rgb = colour


def add_rect(slide, l, t, w, h, colour):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    fill_solid(shape, colour)
    shape.line.fill.background()
    return shape


def text_box(slide, text, l, t, w, h,
             size=18, bold=False, colour=WHITE,
             align=PP_ALIGN.LEFT, wrap=True):
    txBox = slide.shapes.add_textbox(l, t, w, h)
    txBox.word_wrap = wrap
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour
    return txBox


def add_bullets(slide, items, l, t, w, h,
                size=17, colour=WHITE, indent_pt=18):
    txBox = slide.shapes.add_textbox(l, t, w, h)
    txBox.word_wrap = True
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = item
        p.level = 0
        p.space_before = Pt(4)
        p.space_after  = Pt(4)
        run = p.runs[0] if p.runs else p.add_run()
        run.text = item
        run.font.size = Pt(size)
        run.font.color.rgb = colour
        p.bullet = True


def accent_bar(slide, colour=ACCENT):
    add_rect(slide, Inches(0), Inches(7.2), W, Inches(0.3), colour)


# ── TITLE SLIDE ──────────────────────────────────────────────────────────────
def title_slide(prs):
    sl = blank_slide(prs)
    add_rect(sl, 0, 0, W, H, NAVY)
    add_rect(sl, 0, 0, Inches(0.12), H, ACCENT)
    add_rect(sl, 0, Inches(3.1), W, Inches(0.06), BLUE)

    text_box(sl, "A Technology Vision for\nQuantitative Trading",
             Inches(0.5), Inches(0.9), Inches(12), Inches(1.9),
             size=40, bold=True, colour=WHITE)

    text_box(sl, "Thomas Schmelzer",
             Inches(0.5), Inches(3.4), Inches(8), Inches(0.5),
             size=22, bold=True, colour=ACCENT)

    text_box(sl, "May 2026",
             Inches(0.5), Inches(3.95), Inches(8), Inches(0.4),
             size=18, colour=LIGHT_GRAY)

    text_box(sl,
             "Two decades across systematic hedge funds, HFT,\n"
             "family offices and sovereign wealth funds.",
             Inches(0.5), Inches(4.6), Inches(10), Inches(0.8),
             size=16, colour=MID_GRAY)


# ── CONTENT SLIDE HELPER ─────────────────────────────────────────────────────
def content_slide(prs, title, bullets, note=None):
    sl = blank_slide(prs)
    add_rect(sl, 0, 0, W, H, NAVY)
    add_rect(sl, 0, 0, W, Inches(1.15), BLUE)
    add_rect(sl, 0, 0, Inches(0.08), H, ACCENT)
    accent_bar(sl)

    text_box(sl, title,
             Inches(0.3), Inches(0.18), Inches(12.5), Inches(0.85),
             size=28, bold=True, colour=WHITE)

    add_bullets(sl, bullets,
                Inches(0.4), Inches(1.3), Inches(12.5), Inches(5.6),
                size=17, colour=LIGHT_GRAY)

    if note:
        text_box(sl, note,
                 Inches(0.4), Inches(6.85), Inches(12.5), Inches(0.35),
                 size=11, colour=ACCENT)
    return sl


# ── TWO-COLUMN SLIDE ─────────────────────────────────────────────────────────
def two_col_slide(prs, title, left_title, left_items, right_title, right_items):
    sl = blank_slide(prs)
    add_rect(sl, 0, 0, W, H, NAVY)
    add_rect(sl, 0, 0, W, Inches(1.15), BLUE)
    add_rect(sl, 0, 0, Inches(0.08), H, ACCENT)
    accent_bar(sl)
    add_rect(sl, Inches(6.66), Inches(1.2), Inches(0.04), Inches(5.9), BLUE)

    text_box(sl, title,
             Inches(0.3), Inches(0.18), Inches(12.5), Inches(0.85),
             size=28, bold=True, colour=WHITE)

    text_box(sl, left_title,
             Inches(0.4), Inches(1.25), Inches(6), Inches(0.45),
             size=14, bold=True, colour=ACCENT)
    add_bullets(sl, left_items,
                Inches(0.4), Inches(1.75), Inches(6.1), Inches(5.0),
                size=15, colour=LIGHT_GRAY)

    text_box(sl, right_title,
             Inches(6.9), Inches(1.25), Inches(6), Inches(0.45),
             size=14, bold=True, colour=ACCENT)
    add_bullets(sl, right_items,
                Inches(6.9), Inches(1.75), Inches(6.1), Inches(5.0),
                size=15, colour=LIGHT_GRAY)


# ── SECTION DIVIDER ───────────────────────────────────────────────────────────
def divider_slide(prs, label):
    sl = blank_slide(prs)
    add_rect(sl, 0, 0, W, H, NAVY)
    add_rect(sl, 0, Inches(3.4), W, Inches(0.08), ACCENT)
    text_box(sl, label,
             Inches(1), Inches(2.7), Inches(11), Inches(1.2),
             size=38, bold=True, colour=WHITE, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
prs = new_prs()

title_slide(prs)

# ── EXECUTIVE SUMMARY ─────────────────────────────────────────────────────────
content_slide(prs, "Executive Summary", [
    "▸  The old model: researchers in Python/MATLAB, engineers reimplementing in C++",
    "▸  Knowledge fragmented across personal scripts; new hires rebuilt the same tools",
    "▸  Modern ML (PyTorch) made the C++ mandate impossible to sustain",
    "▸  New premise: research and production share the same environment",
    "▸  The same container runs in development and live trading",
    "▸  Team works like a professional kitchen — collaborative, quality shared at every step",
    "▸  Platform exists so the team can focus on the one thing that cannot be bought: the edge",
])

# ── THE PROBLEM ───────────────────────────────────────────────────────────────
divider_slide(prs, "The Problem")

two_col_slide(prs, "The Old Approach & Its Costs",
    "The Script Era",
    [
        "Researchers: mathematicians, physicists, statisticians",
        "Strategies in MATLAB or Python — fast to iterate",
        "Outputs: dense, clever, personal scripts",
        "Logic accumulated over months and years",
        "Code worked but was deeply personal and non-transferable",
    ],
    "Reinventing the Wheel",
    [
        "No shared library — same components rebuilt by every researcher",
        "Diverging results with no clear explanation",
        "Onboarding meant starting from scratch",
        "Institutional knowledge walked out with people",
        "Researcher ambition bounded by programming skill, not ideas",
    ]
)

content_slide(prs, "The Handover Problem", [
    "▸  Strategies passed to C++ engineers for production reimplementation",
    "▸  Handover rarely clean: translation errors, implicit edge cases, precision issues",
    "▸  Research cycles bottlenecked — months between idea and live strategy",
    "▸  PyTorch and modern ML simply cannot be reimplemented in C++",
    "▸  History's verdict: the kitchen was never built",
    "▸  Essential tools reinvented independently by every team — data APIs, portfolio tools, analytics",
    "▸  Result: a sprawling collection of overlapping partial solutions, each owned by whoever wrote it",
    "▸  When things go wrong, accountability diffuses — nobody owns it",
    "▸  Discussed by Marcos Lopez de Prado; put into practice at ADIA's Team Q",
])

# ── A NEW DIRECTION ───────────────────────────────────────────────────────────
divider_slide(prs, "A New Direction")

content_slide(prs, "The Kitchen — A Better Analogy", [
    "▸  Assembly line model fails for knowledge work: interfaces are not fixed in advance",
    "▸  Lopez de Prado's car factory analogy: appealing but misleading",
    "▸  Better image: the professional kitchen — familiar, collaborative, quality at every step",
    "▸  Every station visible to every other; head chef and junior cook share the same pressure",
    "▸  Quality is everyone's responsibility throughout — not inspected at the end",
    "▸  A chef does not build the oven: platform provides the environment",
    "▸  Checkerboard structure: researchers and developers sit together, alternate, collaborate",
    "▸  Boundary between researcher and developer is blurry — by design",
])

content_slide(prs, "Containerization", [
    "▸  Container packages code and the entire runtime: OS libraries, language version, dependencies",
    "▸  Identical environment on laptop, backtesting cluster and live production",
    "▸  Eliminates 'it works on my machine' — every machine runs the same machine",
    "▸  Silent numerical differences between research and production create false confidence",
    "▸  Strategy and environment deployed, tested and promoted together as a versioned artifact",
    "▸  Researcher opens the container and starts — no setup day",
])

# ── BUILDING THE PLATFORM ─────────────────────────────────────────────────────
divider_slide(prs, "Building the Platform")

two_col_slide(prs, "Building the Kitchen — Core Components",
    "Foundation",
    [
        "Data API — clean, reliable, point-in-time",
        "Every strategy, backtest and risk calc depends on it",
        "Researchers ask for what they need; API handles the rest",
        "",
        "Common Strategy Tooling",
        "Portfolio construction, signal combination, position sizing",
        "Implemented once, tested thoroughly, shared by all",
    ],
    "Analytics & Monitoring",
    [
        "Performance Analytics",
        "Returns decomposition, factor attribution, cost accounting",
        "Consistent metrics comparable across strategies and time",
        "",
        "Live Monitoring",
        "Position/exposure tracking, P&L, execution quality",
        "Alerts when behaviour moves outside expected bounds",
    ]
)

content_slide(prs, "Building the Strategies — Rhiza", [
    "▸  Each strategy lives in its own repository",
    "▸  Left unmanaged: CI diverges, Python versions drift, security fixes missed",
    "▸  Rhiza is a living template system — keeps every repo aligned with platform standards",
    "▸  Template changes → pull request in each downstream repo with a clean diff",
    "▸  Owners review, adapt if needed, merge — nothing forced, nothing missed",
    "▸  Already in use at Stanford CVXGRP and Janus Henderson",
    "▸  Kitchen and strategies often built in parallel — but a minimal foundation comes first",
], note="github.com/Jebel-Quant/rhiza")

# ── LIVE TRADING ──────────────────────────────────────────────────────────────
divider_slide(prs, "Live Trading")

two_col_slide(prs, "Live Trading — Proximity & Configuration",
    "Same Artifact, Every Environment",
    [
        "Same container in research, paper trading and production",
        "No reimplementation, no port, no translation step",
        "Strategy code identical across all environments",
        "",
        "Configuration drives the difference",
        "Data source, parameters, risk limits, execution venue",
        "Moving to live = changing config, not rebuilding",
        "Every deployment: known image + known config state",
    ],
    "Prime Broker Connectivity",
    [
        "Strategy expresses intent — execution layer translates",
        "FIX or proprietary API: irrelevant to the strategy",
        "Switching brokers is a configuration change",
        "",
        "Rollback",
        "Rolling back is a config change, not an emergency deployment",
        "Full audit trail: reconstruct exactly what ran and when",
    ]
)

# ── QUALITY & RISK ────────────────────────────────────────────────────────────
divider_slide(prs, "Quality & Risk")

content_slide(prs, "Backtesting — Honesty First", [
    "▸  The gap between backtest and live is rarely bad ideas — it is optimistic backtests",
    "▸  Look-ahead bias: platform enforces strict point-in-time data semantics",
    "▸  Transaction costs: market impact, bid-ask, borrow costs modelled explicitly",
    "▸  Multiple cost scenarios: conservative, realistic, optimistic",
    "▸  Overfitting: walk-forward analysis and out-of-sample evaluation as standard",
    "▸  Culture matters as much as tooling — no tool prevents repeated tweaking",
    "▸  Backtest provides evidence, not certainty; treat strong results with scepticism",
])

content_slide(prs, "Risk Management", [
    "▸  Risk is a layer through the entire platform — not bolted on at the end",
    "▸  Research: factor exposure analysis, stress testing before any live capital",
    "▸  Deployment: pre-trade checks — position, notional and concentration limits",
    "▸  All limits are configuration: versioned, auditable, no code changes needed",
    "▸  Production: continuous monitoring, drawdown limits, automatic halts",
    "▸  Kill switch is a first-class concept — tested regularly, not an emergency procedure",
])

# ── AI & JEBEL QUANT ──────────────────────────────────────────────────────────
divider_slide(prs, "AI & Jebel Quant Research")

content_slide(prs, "The Impact of AI", [
    "▸  AI is used at every stage — building the kitchen, research and production",
    "▸  Platform construction: boilerplate, documentation and code review generated in hours",
    "▸  Research: signal generation, literature review, rapid prototyping",
    "▸  More ideas tested; more reach the stage of serious evaluation",
    "▸  Production: monitoring and anomaly detection — problems surface in minutes",
    "▸  AI removes friction; judgement remains with the team",
    "▸  AI does not know which signals are real — humans carry that responsibility",
])

content_slide(prs, "Jebel Quant Research", [
    "▸  Project infrastructure — Rhiza: living template system, in use at Stanford & Janus Henderson",
    "▸  Data access — normalised, versioned API across asset classes",
    "▸  Portfolio construction — convex optimisation; collaboration with Stephen Boyd & Ron Kahn",
    "▸  Signal combination & analytics — shared vocabulary across the team (jquantstats)",
    "▸  Live trading infrastructure — container model, config management, broker connectivity",
    "",
    "▸  github.com/Jebel-Quant",
], note="All tools designed to be reusable across strategies and organisations")

prs.save("/Users/thomasschmelzer/projects/jqr/repos/platform/vision.pptx")
print("Saved vision.pptx")
